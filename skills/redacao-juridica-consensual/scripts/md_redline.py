#!/usr/bin/env python3
"""Gera DOCX com controle de alterações entre duas versões em Markdown, sem precisar do Word.

Complementa ``word_redline.py`` (comparação nativa do Word) quando o canônico é
Markdown: lê a versão-base e a versão-alvo, alinha os parágrafos por
similaridade e escreve inclusões (``w:ins``) e exclusões (``w:del``) palavra a
palavra. Trechos iguais curtos entre duas alterações são fundidos na marca, para
o redline não virar confete. Depois valida nos dois sentidos: aceitar todas as
alterações reproduz a versão-alvo, e rejeitar todas reproduz a base.

Aceita o Markdown simples (``#`` e ``N. texto``) e o Markdown GENES de
``build_estilo_v2.py`` (``@CAB``, ``@EMENTA``, ``@DATA``, ``@ASS``...), com
``**negrito**``. A numeração dos parágrafos é a automática do Word.

Exemplo::

    md_redline.py v29.md v30.md v29_para_v30_controle_alteracoes.docx --autor Nicholas
"""

from __future__ import annotations

import argparse
import difflib
import re
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
NUM_ID = "910"


# ---------------------------------------------------------------- leitura do Markdown

def linhas(path: Path) -> list[str]:
    return [l.strip() for l in path.read_text(encoding="utf-8").split("\n") if l.strip()]


def formato(linha: str) -> tuple[dict, str | None, str, bool]:
    """(estilo do parágrafo, rótulo em negrito, texto, tudo em negrito)."""
    E = WD_ALIGN_PARAGRAPH
    if linha.startswith("@CAB "):
        corpo = linha[5:]
        if "|" in corpo:
            rot, txt = corpo.split("|", 1)
            return dict(spacing=1.15, after=4), rot.strip(), txt.strip(), False
        return dict(spacing=1.15, after=8), None, corpo, True
    if linha.startswith("@EMENTA "):
        rot, txt = linha[8:].split("|", 1)
        return dict(spacing=1.15, before=4, after=10), rot.strip(), txt.strip(), False
    m = re.match(r"^#{1,6} (.*)$", linha)
    if m:
        return dict(align=E.LEFT, spacing=1.15, before=8, after=4, keep=True), None, m.group(1), True
    for marca, estilo, negrito in (
        ("@DATA ", dict(align=E.LEFT, before=6, after=10), False),
        ("@ASS+ ", dict(align=E.LEFT, spacing=1.15, before=12, after=2), True),
        ("@ASS ", dict(align=E.LEFT, spacing=1.15, after=2), True),
        ("@ASN+ ", dict(align=E.LEFT, spacing=1.15, before=6, after=2), False),
        ("@ASN ", dict(align=E.LEFT, spacing=1.15, after=2), False),
    ):
        if linha.startswith(marca):
            return estilo, None, linha[len(marca):], negrito
    m = re.match(r"^(\d+)\. (.*)$", linha)
    if m:
        return dict(num=True, after=2), None, m.group(2), False
    return dict(), None, linha, False


def texto_visivel(linha: str) -> str:
    _, rot, txt, _ = formato(linha)
    return ((rot + " ") if rot else "") + txt.replace("**", "")


def chave(linha: str) -> str:
    return re.sub(r"^\d+\. ", "", linha)


# ---------------------------------------------------------------- escrita do DOCX

class Redline:
    def __init__(self, autor: str, data: str, fonte: str, corpo_pt: float, funde_iguais: int):
        self.autor, self.data, self.funde_iguais = autor, data, funde_iguais
        self.rev = 0
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(2.0), Cm(3.0), Cm(2.0)
        st = self.d.styles["Normal"]
        st.font.name, st.font.size = fonte, Pt(corpo_pt)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), fonte)
        self.tab = str(int(1.25 * 567))
        self._lista(fonte, corpo_pt)

    def _lista(self, fonte: str, corpo_pt: float) -> None:
        numbering = self.d.part.numbering_part.element
        absn = OxmlElement("w:abstractNum")
        absn.set(qn("w:abstractNumId"), NUM_ID)
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "singleLevel")
        absn.append(mlt)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        for tag, val in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1."), ("w:lvlJc", "left"), ("w:suff", "tab")):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            lvl.append(e)
        rp = OxmlElement("w:rPr")
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            f.set(qn(a), fonte)
        rp.append(f)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(corpo_pt * 2)))
        rp.append(sz)
        lvl.append(rp)
        absn.append(lvl)
        primeiro = numbering.find(qn("w:num"))
        primeiro.addprevious(absn) if primeiro is not None else numbering.append(absn)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), NUM_ID)
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), NUM_ID)
        num.append(ref)
        numbering.append(num)

    def marca(self, e):
        self.rev += 1
        e.set(qn("w:id"), str(self.rev))
        e.set(qn("w:author"), self.autor)
        e.set(qn("w:date"), self.data)
        return e

    def paragrafo(self, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.5, before=0, after=6, keep=False, num=False, revisao=None):
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.alignment, pf.line_spacing, pf.space_before, pf.space_after = align, spacing, Pt(before), Pt(after)
        if keep:
            pf.keep_with_next = True
        if num:
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), "0")
            ni = OxmlElement("w:numId")
            ni.set(qn("w:val"), NUM_ID)
            numPr.append(il)
            numPr.append(ni)
            pPr.insert(0, numPr)
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "left")
            tab.set(qn("w:pos"), self.tab)
            tabs.append(tab)
            numPr.addnext(tabs)
        if revisao:  # marca de parágrafo inteiro incluído ou excluído
            pPr = p._p.get_or_add_pPr()
            rPr = pPr.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                pPr.append(rPr)
            rPr.append(self.marca(OxmlElement("w:" + revisao)))
        return p

    @staticmethod
    def _run(texto: str, negrito: bool, excluido: bool = False):
        r = OxmlElement("w:r")
        if negrito:
            rp = OxmlElement("w:rPr")
            rp.append(OxmlElement("w:b"))
            r.append(rp)
        t = OxmlElement("w:delText" if excluido else "w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = texto
        r.append(t)
        return r

    @staticmethod
    def tokens(texto: str, tudo_negrito: bool = False) -> list[tuple[str, bool]]:
        out = []
        for i, parte in enumerate(re.split(r"\*\*", texto)):
            b = tudo_negrito or i % 2 == 1
            out += [(tk, b) for tk in re.findall(r"\s+|[^\s]+", parte)]
        return out

    def emite(self, p, toks, modo: str) -> None:
        grupos: list[tuple[list[str], bool]] = []
        for tk, b in toks:
            if grupos and grupos[-1][1] == b:
                grupos[-1][0].append(tk)
            else:
                grupos.append(([tk], b))
        for tks, b in grupos:
            txt = "".join(tks)
            if modo == "eq":
                p._p.append(self._run(txt, b))
            else:
                w = self.marca(OxmlElement("w:ins" if modo == "ins" else "w:del"))
                w.append(self._run(txt, b, excluido=modo == "del"))
                p._p.append(w)

    def diff(self, p, a: str, b: str, tudo_negrito: bool) -> None:
        ta, tb = self.tokens(a, tudo_negrito), self.tokens(b, tudo_negrito)
        ops = difflib.SequenceMatcher(None, ta, tb, autojunk=False).get_opcodes()
        fundidos: list[tuple] = []
        for k, (op, i1, i2, j1, j2) in enumerate(ops):
            curto = op == "equal" and (i2 - i1) <= self.funde_iguais and 0 < k < len(ops) - 1
            if curto or op != "equal":
                if fundidos and fundidos[-1][0] == "chg":
                    _, a1, _, b1, _ = fundidos[-1]
                    fundidos[-1] = ("chg", a1, i2, b1, j2)
                else:
                    fundidos.append(("chg", i1, i2, j1, j2))
            else:
                fundidos.append((op, i1, i2, j1, j2))
        for op, i1, i2, j1, j2 in fundidos:
            if op == "equal":
                self.emite(p, ta[i1:i2], "eq")
            else:
                if i2 > i1:
                    self.emite(p, ta[i1:i2], "del")
                if j2 > j1:
                    self.emite(p, tb[j1:j2], "ins")


def plano_alinhamento(A: list[str], B: list[str], limiar: float = 0.3) -> list[tuple[str, int | None, int | None]]:
    ka, kb = [chave(l) for l in A], [chave(l) for l in B]

    def alinhar(ai: list[int], bi: list[int]):
        n, m = len(ai), len(bi)
        sim = [[difflib.SequenceMatcher(None, ka[ai[i]], kb[bi[j]], autojunk=False).ratio() for j in range(m)] for i in range(n)]
        best = [[0.0] * (m + 1) for _ in range(n + 1)]
        trace: list[list[str | None]] = [[None] * (m + 1) for _ in range(n + 1)]
        for i in range(n + 1):
            for j in range(m + 1):
                if i == 0 and j == 0:
                    continue
                opts = []
                if i and j and sim[i - 1][j - 1] >= limiar:
                    opts.append((best[i - 1][j - 1] + sim[i - 1][j - 1], "par"))
                if i:
                    opts.append((best[i - 1][j], "del"))
                if j:
                    opts.append((best[i][j - 1], "ins"))
                best[i][j], trace[i][j] = max(opts)
        i, j, seq = n, m, []
        while i or j:
            t = trace[i][j]
            if t == "par":
                seq.append(("par", ai[i - 1], bi[j - 1]))
                i, j = i - 1, j - 1
            elif t == "del":
                seq.append(("del", ai[i - 1], None))
                i -= 1
            else:
                seq.append(("ins", None, bi[j - 1]))
                j -= 1
        return seq[::-1]

    plano: list[tuple[str, int | None, int | None]] = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, ka, kb, autojunk=False).get_opcodes():
        if op == "equal":
            plano += [("eq", i, j) for i, j in zip(range(i1, i2), range(j1, j2))]
        else:
            plano += alinhar(list(range(i1, i2)), list(range(j1, j2)))
    return plano


def gerar(base: Path, alvo: Path, saida: Path, autor: str, data: str | None = None, fonte: str = "Times New Roman",
          corpo_pt: float = 10.5, funde_iguais: int = 3) -> dict[str, int]:
    A, B = linhas(base), linhas(alvo)
    rl = Redline(autor, data or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), fonte, corpo_pt, funde_iguais)
    stats = {"eq": 0, "par": 0, "ins": 0, "del": 0}
    for op, i, j in plano_alinhamento(A, B):
        stats[op] += 1
        linha = B[j] if j is not None else A[i]  # type: ignore[index]
        estilo, rot, _, negrito = formato(linha)
        p = rl.paragrafo(**estilo, revisao={"ins": "ins", "del": "del"}.get(op))
        if rot:
            r = p.add_run(rot + " ")
            r.bold = True
        if op == "eq":
            rl.emite(p, rl.tokens(formato(B[j])[2], negrito), "eq")  # type: ignore[index]
        elif op == "par":
            rl.diff(p, formato(A[i])[2], formato(B[j])[2], negrito)  # type: ignore[index]
        elif op == "ins":
            rl.emite(p, rl.tokens(formato(B[j])[2], negrito), "ins")  # type: ignore[index]
        else:
            rl.emite(p, rl.tokens(formato(A[i])[2], negrito), "del")  # type: ignore[index]
    rl.d.settings.element.insert(0, OxmlElement("w:trackRevisions"))
    rl.d.save(saida)
    return {**stats, "revisoes": rl.rev}


# ---------------------------------------------------------------- validação bidirecional

def textos_resolvidos(docx: Path, modo: str) -> list[str]:
    root = etree.fromstring(zipfile.ZipFile(docx).read("word/document.xml"))
    out = []
    for p in root.iter(W + "p"):
        pr = p.find(W + "pPr")
        marca = None
        if pr is not None and pr.find(W + "rPr") is not None:
            r = pr.find(W + "rPr")
            marca = "ins" if r.find(W + "ins") is not None else "del" if r.find(W + "del") is not None else None
        if (modo == "aceitar" and marca == "del") or (modo == "rejeitar" and marca == "ins"):
            continue
        partes = []
        for el in p.iter():
            if el.tag not in (W + "t", W + "delText"):
                continue
            anc = {a.tag for a in el.iterancestors()}
            if (modo == "aceitar" and W + "del" in anc) or (modo == "rejeitar" and W + "ins" in anc):
                continue
            partes.append(el.text or "")
        txt = "".join(partes)
        if txt:
            out.append(txt)
    return out


def validar(docx: Path, base: Path, alvo: Path) -> dict[str, bool]:
    return {
        "aceitar_reproduz_alvo": textos_resolvidos(docx, "aceitar") == [texto_visivel(l) for l in linhas(alvo)],
        "rejeitar_reproduz_base": textos_resolvidos(docx, "rejeitar") == [texto_visivel(l) for l in linhas(base)],
    }


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("base", type=Path)
    ap.add_argument("alvo", type=Path)
    ap.add_argument("saida", type=Path)
    ap.add_argument("--autor", required=True, help="nome que o Word mostra nas revisões")
    ap.add_argument("--data", help="ISO 8601 UTC; padrão: agora")
    ap.add_argument("--fonte", default="Times New Roman")
    ap.add_argument("--corpo-pt", type=float, default=10.5)
    ap.add_argument("--funde-iguais", type=int, default=3,
                    help="funde na marca trechos iguais de até N tokens entre duas alterações (0 desliga)")
    args = ap.parse_args(argv)
    if not args.autor.strip():
        ap.error("--autor não pode ser vazio")
    stats = gerar(args.base, args.alvo, args.saida, args.autor.strip(), args.data, args.fonte, args.corpo_pt, args.funde_iguais)
    print(f"parágrafos iguais {stats['eq']} | alterados {stats['par']} | incluídos {stats['ins']} | "
          f"excluídos {stats['del']} | revisões {stats['revisoes']}")
    v = validar(args.saida, args.base, args.alvo)
    print(f"aceitar todas reproduz o alvo: {v['aceitar_reproduz_alvo']} | rejeitar todas reproduz a base: {v['rejeitar_reproduz_base']}")
    return 0 if all(v.values()) else 1


if __name__ == "__main__":
    sys.exit(main())
