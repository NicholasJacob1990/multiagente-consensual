"""Gera o parecer no perfil GENES: cabeçalho, ementa, seções e parágrafos com
numeração automática do Word — número alinhado à margem esquerda e tabulação
até RECUO_NUM antes da primeira linha; as linhas seguintes voltam à margem."""
import sys, re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CORPO_PT = 10.5
RECUO_NUM_CM = 1.25
ESPACO_DEPOIS_PT = 2   # entre parágrafos numerados do corpo

src, out = sys.argv[1], sys.argv[2]
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.0); sec.top_margin = Cm(3.0); sec.bottom_margin = Cm(2.0)
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(CORPO_PT)
rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
if rf is None:
    rf = rpr.makeelement(qn('w:rFonts'), {}); rpr.append(rf)
for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), 'Times New Roman')

TW = str(int(RECUO_NUM_CM * 567))

def garante_lista():
    numbering = d.part.numbering_part.element
    absn = OxmlElement('w:abstractNum'); absn.set(qn('w:abstractNumId'), '910')
    mlt = OxmlElement('w:multiLevelType'); mlt.set(qn('w:val'), 'singleLevel'); absn.append(mlt)
    lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), '0')
    for tag, val in (('w:start', '1'), ('w:numFmt', 'decimal'), ('w:lvlText', '%1.'), ('w:lvlJc', 'left'), ('w:suff', 'tab')):
        e = OxmlElement(tag); e.set(qn('w:val'), val); lvl.append(e)
    pPr = OxmlElement('w:pPr')
    tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab'); tab.set(qn('w:val'), 'num'); tab.set(qn('w:pos'), TW); tabs.append(tab); pPr.append(tabs)
    ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '0'); ind.set(qn('w:firstLine'), '0'); pPr.append(ind)
    lvl.append(pPr)
    rp = OxmlElement('w:rPr'); f = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'): f.set(qn(a), 'Times New Roman')
    rp.append(f); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(CORPO_PT * 2))); rp.append(sz); lvl.append(rp)
    absn.append(lvl)
    first_num = numbering.find(qn('w:num'))
    (first_num.addprevious(absn) if first_num is not None else numbering.append(absn))
    num = OxmlElement('w:num'); num.set(qn('w:numId'), '910')
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), '910'); num.append(ref)
    numbering.append(num)
    return '910'

NUM_ID = garante_lista()

def numerar(p):
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement('w:numPr')
    il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0'); numPr.append(il)
    ni = OxmlElement('w:numId'); ni.set(qn('w:val'), NUM_ID); numPr.append(ni)
    pStyle = pPr.find(qn('w:pStyle'))
    (pStyle.addnext(numPr) if pStyle is not None else pPr.insert(0, numPr))
    tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab'); tab.set(qn('w:val'), 'left'); tab.set(qn('w:pos'), TW); tabs.append(tab)
    numPr.addnext(tabs)

def para(text='', bold_label=None, bold_all=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.5, before=0, after=6, keep=False, num=False):
    p = d.add_paragraph()
    pf = p.paragraph_format; pf.alignment = align; pf.line_spacing = spacing
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if keep: pf.keep_with_next = True
    if num:
        numerar(p); pf.left_indent = Cm(0); pf.first_line_indent = Cm(0)
    if bold_label:
        r = p.add_run(bold_label + ' '); r.bold = True
    if text:
        r = p.add_run(text); r.bold = bold_all
    return p

esperado = 0
for raw in open(src, encoding='utf-8').read().split('\n'):
    line = raw.strip()
    if not line: continue
    if line.startswith('@CAB '):
        body = line[5:]
        if '|' in body:
            lab, txt = body.split('|', 1); para(txt.strip(), bold_label=lab.strip(), spacing=1.15, after=4)
        else:
            para(body, bold_all=True, spacing=1.15, after=8)
    elif line.startswith('@EMENTA '):
        lab, txt = line[8:].split('|', 1); para(txt.strip(), bold_label=lab.strip(), spacing=1.15, before=4, after=10)
    elif line.startswith('## '):
        para(line[3:], bold_all=True, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.15, before=8, after=4, keep=True)
    elif line.startswith('@DATA '):
        para(line[6:], align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=10)
    elif line.startswith('@ASS '):
        para(line[5:], bold_all=True, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.15, after=2)
    else:
        m = re.match(r'^(\d+)\. (.*)$', line)
        if not m:
            raise SystemExit(f'linha de corpo sem número: {line[:60]}')
        esperado += 1
        if int(m.group(1)) != esperado:
            raise SystemExit(f'numeração fora de sequência: {m.group(1)} (esperado {esperado})')
        para(m.group(2), num=True, after=ESPACO_DEPOIS_PT)
d.save(out); print('ok', esperado, 'parágrafos numerados automaticamente')
