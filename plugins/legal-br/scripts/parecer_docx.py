#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""parecer_docx.py — entrega de parecer em .docx no padrão institucional.

Resolve o que o Markdown não resolve: fonte/margens/espaçamento da casa,
numeração AUTOMÁTICA de parágrafos (a do Word, não "1." digitado), notas de
rodapé e uma auditoria que pega os defeitos que revisor humano cobra.

Subcomandos
-----------
  check                        dependências
  converter  ENTRADA.md SAIDA.docx     markdown -> docx já no padrão
  formatar   ARQ.docx          fonte, margens, espaçamento, justificado
  numerar    ARQ.docx          numeração automática dos parágrafos de corpo
  nota       ARQ.docx --par N --texto "..."      nota de rodapé
  auditar    ARQ.docx          relatório JSON dos defeitos típicos
  ementa     ARQ.docx [--max 80] [--nova "..."]  confere/substitui a ementa

Toda escrita cria backup `<arquivo>.bak-<AAAAMMDD-HHMM>.docx` ao lado.
"""
from __future__ import annotations

import argparse, copy, json, re, shutil, subprocess, sys, zipfile
from datetime import datetime
from pathlib import Path

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    docx = None

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"

# Padrão da casa (GENSU/GENES - CODEMGE). Ajuste aqui para outra instituição.
PADRAO = dict(fonte="Times New Roman", corpo=12, nota=10, entrelinha=1.5,
              margens=(3.0, 2.0, 3.0, 2.0), espaco_depois=6, recuo_num=1.25)

# Parágrafos que NUNCA recebem número: cabeçalho, ementa, títulos e fecho.
NAO_NUMERA = re.compile(
    r"^(PARECER\b|Processo\s+n|Interessad|EMENTA:|REFER[ÊE]NCIAS\b"
    r"|[IVXLC]+\s*[—\-–]\s*\S"                    # "IV — DA NATUREZA..."
    r"|É o parecer|É o que me parece|Ante o exposto|Antes da celebra"
    r"|Advogad|Ger[êe]ncia\b"
    # cabeçalho, seções e fecho do perfil GENES (references/perfil-parecer-genes.md)
    r"|PROCESSO\s+SEI|ASSUNTO:|CONSULENTE:|[ÓO]RG[ÃA]O\s+CONSULTIVO|EMENTA\s+CONSULTIVA"
    r"|[IVXLC]+\.\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ]|Belo\s+Horizonte,|Diretoria\s+Jur[íi]dica)", re.I)

FECHO = re.compile(r"^(É o parecer|É o que me parece|S\.M\.J\.)", re.I)

TYPOS = {  # colhidos em revisão real; francesismos e trocas de teclado
    "Entranto": "Entretanto", "eventualemnte": "eventualmente",
    "relative ": "relativo ", "a confirmer": "a confirmar",
    "hiposuficiente": "hipossuficiente", "atravez": "através",
    "excessão": "exceção", "previlégio": "privilégio",
}

# Citação de precedente sem relator/órgão/data — o defeito nº 1 de referências.
PRECEDENTE = re.compile(r"\b(REsp|AREsp|AgInt|RE|ARE|ADI|ADC|ADPF|RMS|HC|MS)\b[^.;)]*?"
                        r"n?[ºo]?\s*[\d][\d.\-/]{4,}", re.I)
TEM_AUTORIA = re.compile(r"\b(Rel\.|Relator|Relatora|Turma|Plen[áa]rio|Se[çc][ãa]o|"
                         r"DJe|DJEN|julgad|j\.\s*\d)", re.I)


def _agora() -> str:
    return datetime.now().strftime("%Y%m%d-%H%M")


def _backup(p: Path) -> Path:
    b = p.with_suffix(f".bak-{_agora()}.docx")
    shutil.copy2(p, b)
    return b


def _norm(s: str) -> str:
    """NBSP quebra qualquer match de texto vindo do Word."""
    return s.replace("\xa0", " ").replace("‑", "-")


def _rpr(par):
    if not par.runs:
        return None
    r = par.runs[0]._r.find(qn("w:rPr"))
    return copy.deepcopy(r) if r is not None else None


def _mk_run(rPr, texto):
    r = OxmlElement("w:r")
    if rPr is not None:
        r.append(copy.deepcopy(rPr))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = texto
    r.append(t)
    return r


def set_texto(par, texto):
    """Reescreve o parágrafo num único run, preservando o rPr do primeiro
    (é o que mantém o negrito da ementa e a fonte herdada)."""
    rPr = _rpr(par)
    el = par._p
    for c in list(el):
        if c.tag != qn("w:pPr"):
            el.remove(c)
    el.append(_mk_run(rPr, texto))


def inserir_depois(par, textos):
    """Insere parágrafos clonando o pPr do de referência (herda numeração,
    estilo e espaçamento). Sem isso o parágrafo novo sai fora do padrão."""
    ref, rPr, ancora = par._p, _rpr(par), par._p
    for txt in textos:
        novo = copy.deepcopy(ref)
        for a in (W14 + "paraId", W14 + "textId"):
            novo.attrib.pop(a, None)
        for c in list(novo):
            if c.tag != qn("w:pPr"):
                novo.remove(c)
        novo.append(_mk_run(rPr, txt))
        ancora.addnext(novo)
        ancora = novo


def tem_numeracao(par) -> bool:
    # ARMADILHA: numPr fica dentro de pPr. find() direto no w:p devolve None
    # sempre e faz a auditoria mentir "0 numerados".
    return par._p.find(".//" + W + "numPr") is not None


# ---------------------------------------------------------------- numeração

def _garante_lista(doc, num_id="900") -> str:
    """Cria abstractNum+num em numbering.xml (decimal, número na margem e tab)."""
    numbering = doc.part.numbering_part.element
    for n in numbering.findall(qn("w:num")):
        if n.get(qn("w:numId")) == num_id:
            return num_id
    aid = "900"
    absn = OxmlElement("w:abstractNum")
    absn.set(qn("w:abstractNumId"), aid)
    mlt = OxmlElement("w:multiLevelType"); mlt.set(qn("w:val"), "singleLevel")
    absn.append(mlt)
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "decimal"); lvl.append(fmt)
    txt = OxmlElement("w:lvlText"); txt.set(qn("w:val"), "%1."); lvl.append(txt)
    jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); lvl.append(jc)
    suff = OxmlElement("w:suff"); suff.set(qn("w:val"), "tab"); lvl.append(suff)
    pPr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "0")                                   # número na margem
    ind.set(qn("w:firstLine"), str(int(PADRAO["recuo_num"] * 567)))  # texto recuado
    pPr.append(ind); lvl.append(pPr)
    absn.append(lvl)
    numbering.insert(0, absn)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), num_id)
    ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), aid)
    num.append(ref); numbering.append(num)
    return num_id


def aplicar_numeracao(par, num_id="900"):
    pPr = par._p.get_or_add_pPr()
    if pPr.find(qn("w:numPr")) is not None:
        return False
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), num_id)
    numPr.append(ilvl); numPr.append(nid)
    # ordem do schema CT_PPr: pStyle vem antes de numPr, tabs depois
    pStyle = pPr.find(qn("w:pStyle"))
    (pStyle.addnext(numPr) if pStyle is not None else pPr.insert(0, numPr))
    if pPr.find(qn("w:tabs")) is None:
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "left")
        tab.set(qn("w:pos"), str(int(PADRAO["recuo_num"] * 567)))
        tabs.append(tab); numPr.addnext(tabs)
    return True


def cmd_numerar(a):
    doc = docx.Document(a.arquivo)
    num_id = _garante_lista(doc)
    n, no_fecho = 0, False
    for p in doc.paragraphs:
        t = _norm(p.text).strip()
        # entre "É o parecer" e "REFERÊNCIAS" só há assinatura, cargo e unidade
        if FECHO.match(t):
            no_fecho = True
        elif t.upper().startswith("REFER"):
            no_fecho = False
        if not t or no_fecho or NAO_NUMERA.match(t) or tem_numeracao(p):
            continue
        if aplicar_numeracao(p, num_id):
            n += 1
    bkp = _backup(Path(a.arquivo)); doc.save(a.arquivo)
    print(json.dumps({"numerados_agora": n, "backup": bkp.name}, ensure_ascii=False))


# ---------------------------------------------------------------- formatação

def cmd_formatar(a):
    doc = docx.Document(a.arquivo)
    for s in doc.sections:
        t, r, b, l = PADRAO["margens"]
        s.top_margin, s.right_margin = Cm(t), Cm(r)
        s.bottom_margin, s.left_margin = Cm(b), Cm(l)
    normal = doc.styles["Normal"]
    normal.font.name = PADRAO["fonte"]
    normal.font.size = Pt(PADRAO["corpo"])
    rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rpr.set(qn(attr), PADRAO["fonte"])
    for p in doc.paragraphs:
        if not _norm(p.text).strip():
            continue
        pf = p.paragraph_format
        pf.line_spacing = PADRAO["entrelinha"]
        pf.space_after = Pt(PADRAO["espaco_depois"])
        if p.alignment is None:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.name = PADRAO["fonte"]
            if r.font.size is None:
                r.font.size = Pt(PADRAO["corpo"])
    bkp = _backup(Path(a.arquivo)); doc.save(a.arquivo)
    print(json.dumps({"formatado": True, "padrao": PADRAO["fonte"],
                      "backup": bkp.name}, ensure_ascii=False))


# ---------------------------------------------------------------- nota

def cmd_nota(a):
    """Nota de rodapé. python-docx não expõe API para isso; editamos
    footnotes.xml dentro do zip.

    ARMADILHA: alterar o objeto XML da part em memória NÃO persiste no save() —
    a gravação tem de reescrever o zip. Estratégia: marcar o ponto de inserção
    com um token via python-docx, salvar, e então trocar o token pela referência
    já no XML bruto."""
    p = Path(a.arquivo)
    with zipfile.ZipFile(p) as z:
        if "word/footnotes.xml" not in z.namelist():
            sys.exit("documento sem footnotes.xml — insira uma nota pelo Word uma vez "
                     "para o Word criar a part, depois use este comando")
        fn = z.read("word/footnotes.xml").decode("utf-8")

    ids = [int(m) for m in re.findall(r'<w:footnote[^>]*w:id="(-?\d+)"', fn)]
    novo_id = max(ids) + 1 if ids else 1

    bkp = _backup(p)
    doc = docx.Document(str(p))
    try:
        alvo = doc.paragraphs[a.par]
    except IndexError:
        sys.exit(f"parágrafo {a.par} não existe (documento tem {len(doc.paragraphs)})")
    marca = f"@@NOTA{novo_id}@@"
    set_texto(alvo, _norm(alvo.text).rstrip() + marca)
    doc.save(str(p))

    esc = (a.texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    nota = (f'<w:footnote w:id="{novo_id}"><w:p><w:pPr>'
            f'<w:pStyle w:val="Textodenotaderodap"/></w:pPr>'
            f'<w:r><w:rPr><w:rStyle w:val="Refdenotaderodap"/></w:rPr><w:footnoteRef/></w:r>'
            f'<w:r><w:t xml:space="preserve"> {esc}</w:t></w:r></w:p></w:footnote>')
    ref = ('<w:r><w:rPr><w:rStyle w:val="Refdenotaderodap"/></w:rPr>'
           f'<w:footnoteReference w:id="{novo_id}"/></w:r>')

    with zipfile.ZipFile(p) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    dx = partes["word/document.xml"].decode("utf-8")
    if marca not in dx:
        shutil.copy2(bkp, p)
        sys.exit("marcador não encontrado no XML — nada alterado (backup restaurado)")
    # fecha o run que contém a marca e emenda a referência logo depois dele
    i = dx.index(marca)
    fim_run = dx.index("</w:r>", i) + len("</w:r>")
    dx = dx[:fim_run] + ref + dx[fim_run:]
    dx = dx.replace(marca, "")
    partes["word/document.xml"] = dx.encode("utf-8")
    partes["word/footnotes.xml"] = fn.replace(
        "</w:footnotes>", nota + "</w:footnotes>").encode("utf-8")

    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, dados in partes.items():
            z.writestr(nome, dados)
    print(json.dumps({"nota_id": novo_id, "paragrafo": a.par,
                      "backup": bkp.name}, ensure_ascii=False))


# ---------------------------------------------------------------- auditoria

def cmd_auditar(a):
    doc = docx.Document(a.arquivo)
    pars = doc.paragraphs
    achados, sem_num, refs_incompletas = [], [], []
    ementa = None
    for i, p in enumerate(pars):
        t = _norm(p.text).strip()
        if not t:
            continue
        if t.startswith("EMENTA:") and ementa is None:
            ementa = (i, t)
        for errado, certo in TYPOS.items():
            if errado in t:
                achados.append({"tipo": "typo", "par": i,
                                "achado": errado, "sugestao": certo})
        if t.endswith(","):
            achados.append({"tipo": "frase_truncada", "par": i, "trecho": t[-60:]})
        if not tem_numeracao(p) and not NAO_NUMERA.match(t):
            sem_num.append({"par": i, "trecho": t[:70]})
        if PRECEDENTE.search(t) and not TEM_AUTORIA.search(t) and len(t) < 200:
            refs_incompletas.append({"par": i, "trecho": t[:110]})

    with zipfile.ZipFile(a.arquivo) as z:
        dx = z.read("word/document.xml").decode("utf-8")
    notas_corpo = len(re.findall(r"<w:footnoteReference", dx))

    rel = {
        "paragrafos": len(pars),
        "numerados": sum(1 for p in pars if tem_numeracao(p)),
        "sem_numeracao_suspeitos": sem_num,
        "ementa_palavras": len(ementa[1].split()) if ementa else None,
        "ementa_excede": bool(ementa and len(ementa[1].split()) > a.max_ementa),
        "notas_rodape_no_corpo": notas_corpo,
        "citacoes_sem_relator_ou_data": refs_incompletas,
        "defeitos_texto": achados,
    }
    rel["ok"] = not (rel["ementa_excede"] or achados or refs_incompletas or sem_num)
    print(json.dumps(rel, ensure_ascii=False, indent=2))


# ---------------------------------------------------------------- ementa

def cmd_ementa(a):
    doc = docx.Document(a.arquivo)
    alvo = next((p for p in doc.paragraphs[:15]
                 if _norm(p.text).strip().startswith("EMENTA:")), None)
    if alvo is None:
        sys.exit("ementa não localizada (procure um parágrafo iniciado por 'EMENTA:')")
    atual = _norm(alvo.text).strip()
    if not a.nova:
        print(json.dumps({"palavras": len(atual.split()), "max": a.max,
                          "excede": len(atual.split()) > a.max, "texto": atual},
                         ensure_ascii=False, indent=2))
        return
    if len(a.nova.split()) > a.max:
        sys.exit(f"a nova ementa tem {len(a.nova.split())} palavras (máx {a.max})")
    bkp = _backup(Path(a.arquivo))
    set_texto(alvo, a.nova)            # preserva o negrito
    doc.save(a.arquivo)
    print(json.dumps({"antes": len(atual.split()), "depois": len(a.nova.split()),
                      "backup": bkp.name}, ensure_ascii=False))


# ---------------------------------------------------------------- converter

def cmd_converter(a):
    if not shutil.which("pandoc"):
        sys.exit("pandoc não encontrado: brew install pandoc")
    subprocess.run(["pandoc", a.entrada, "-o", a.saida], check=True)
    for fn, args in ((cmd_formatar, {"arquivo": a.saida}),
                     (cmd_numerar, {"arquivo": a.saida})):
        fn(argparse.Namespace(**args))
    print(json.dumps({"gerado": a.saida}, ensure_ascii=False))


def cmd_check(_):
    print(json.dumps({
        "python_docx": docx is not None,
        "pandoc": bool(shutil.which("pandoc")),
        "padrao": PADRAO,
    }, ensure_ascii=False, indent=2))


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("check").set_defaults(func=cmd_check)

    c = sub.add_parser("converter"); c.add_argument("entrada"); c.add_argument("saida")
    c.set_defaults(func=cmd_converter)

    for nome, fn in (("formatar", cmd_formatar), ("numerar", cmd_numerar)):
        s = sub.add_parser(nome); s.add_argument("arquivo"); s.set_defaults(func=fn)

    n = sub.add_parser("nota"); n.add_argument("arquivo")
    n.add_argument("--par", type=int, required=True, help="índice do parágrafo")
    n.add_argument("--texto", required=True); n.set_defaults(func=cmd_nota)

    au = sub.add_parser("auditar"); au.add_argument("arquivo")
    au.add_argument("--max-ementa", type=int, default=80, dest="max_ementa")
    au.set_defaults(func=cmd_auditar)

    e = sub.add_parser("ementa"); e.add_argument("arquivo")
    e.add_argument("--max", type=int, default=80); e.add_argument("--nova")
    e.set_defaults(func=cmd_ementa)

    a = ap.parse_args()
    if docx is None and a.cmd != "check":
        sys.exit("instale python-docx: python3 -m pip install --user python-docx")
    a.func(a)


if __name__ == "__main__":
    main()
